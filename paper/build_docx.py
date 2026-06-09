"""build_docx.py — render paper/PAPER.md to a clean, Google-Docs-importable paper.docx
(embedded charts, tables, headings). Upload to Drive -> opens as a native Google Doc.
Purpose-built Markdown subset parser for this paper's constructs."""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(__file__)
FIG = os.path.join(HERE, "figures")
INK = RGBColor(0x15, 0x20, 0x2b); MUTE = RGBColor(0x5b, 0x6b, 0x7b)
BLUE = RGBColor(0x2f, 0x6d, 0xf0)

doc = Document()
# base style: clean sans, Google-Docs-ish
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11); st.font.color.rgb = INK
for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
    try: doc.styles[s].font.name = "Arial"
    except Exception: pass

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[.+?\]\(.+?\))")
def add_runs(p, text):
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        elif tok.startswith("[") and "](" in tok:
            label = tok[1:tok.index("](")]; r = p.add_run(label); r.font.color.rgb = BLUE
        else:
            p.add_run(tok)

def emit_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for i, r in enumerate(cells) if not (i == 1 and set("".join(r)) <= set("-: "))]
    if not cells: return
    t = doc.add_table(rows=len(cells), cols=len(cells[0])); t.style = "Light Grid Accent 1"
    for ri, row in enumerate(cells):
        for ci, val in enumerate(row):
            if ci < len(t.rows[ri].cells):
                cell = t.rows[ri].cells[ci]; cell.text = ""
                add_runs(cell.paragraphs[0], val)
                if ri == 0:
                    for rn in cell.paragraphs[0].runs: rn.bold = True
    doc.add_paragraph()

md = open(os.path.join(HERE, "PAPER.md")).read().splitlines()

# Title + hero chart up front
title = md[0].lstrip("# ").strip()
h = doc.add_heading(title, level=0)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Technical report  ·  model: huggingface.co/dennisonb/reversible-circuit-8b-tool  ·  code: github.com/dennisonbertram/reversible-circuit-llm")
r.font.size = Pt(9); r.font.color.rgb = MUTE
doc.add_picture(os.path.join(FIG, "fig0_hero.png"), width=Inches(6.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

i, tbl = 1, []
def flush_table():
    global tbl
    if tbl: emit_table(tbl); tbl = []

while i < len(md):
    ln = md[i].rstrip()
    if ln.lstrip().startswith("|") and "|" in ln[1:]:
        tbl.append(ln); i += 1; continue
    flush_table()
    img = re.match(r"!\[(.*?)\]\((.+?)\)", ln.strip())
    if img:
        path = img.group(2)
        fp = os.path.join(HERE, path) if not os.path.isabs(path) else path
        if os.path.exists(fp):
            doc.add_picture(fp, width=Inches(5.8)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(); cr = cap.add_run(img.group(1)); cr.italic = True
            cr.font.size = Pt(9); cr.font.color.rgb = MUTE; cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1; continue
    if ln.startswith("### "): doc.add_heading(ln[4:].strip(), level=3)
    elif ln.startswith("## "): doc.add_heading(ln[3:].strip(), level=2)
    elif ln.startswith("# "): doc.add_heading(ln[2:].strip(), level=1)
    elif ln.strip() == "---": pass
    elif ln.lstrip().startswith(("- ", "* ", "> - ", "> * ")):
        txt = ln.lstrip().lstrip(">").strip()[2:]
        add_runs(doc.add_paragraph(style="List Bullet"), txt)
    elif ln.startswith(">"):
        txt = ln.lstrip(">").strip()
        if txt:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, txt)
            for rn in p.runs: rn.font.color.rgb = MUTE
    elif ln.strip() == "":
        pass
    else:
        add_runs(doc.add_paragraph(), ln)
    i += 1
flush_table()

out = os.path.join(HERE, "paper.docx")
doc.save(out)
print("wrote", out, "(", round(os.path.getsize(out)/1024), "KB )")
